from typing import Tuple
from io import BytesIO
from datetime import datetime
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage
from reportlab.lib.units import inch
from reportlab.lib import colors
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pptx import Presentation
from pptx.util import Inches as PPTInches, Pt as PPTPt
from pptx.enum.text import PP_ALIGN
import csv
from PIL import Image as PILImage


class ExportService:
    """Serviço para exportar documentos em múltiplos formatos"""

    @staticmethod
    def export_pdf(
        title: str,
        content: str,
        logo_path: str = None
    ) -> Tuple[BytesIO, str]:
        """Exporta documento para PDF"""
        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            topMargin=0.5*inch,
            bottomMargin=0.5*inch,
        )

        story = []
        styles = getSampleStyleSheet()

        # Adicionar logo se fornecido
        if logo_path:
            try:
                img = RLImage(logo_path, width=1*inch, height=0.5*inch)
                story.append(img)
                story.append(Spacer(1, 0.3*inch))
            except Exception as e:
                print(f"Erro ao inserir logo: {e}")

        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1f2937'),
            spaceAfter=30,
        )
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 0.2*inch))

        # Conteúdo
        paragraphs = content.split('\n')
        for para in paragraphs:
            if para.strip():
                story.append(Paragraph(para, styles['Normal']))
            else:
                story.append(Spacer(1, 0.1*inch))

        # Rodapé com data
        story.append(Spacer(1, 0.5*inch))
        footer_text = f"Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M:%S')}"
        story.append(Paragraph(footer_text, styles['Normal']))

        doc.build(story)
        buffer.seek(0)

        return buffer, "application/pdf"

    @staticmethod
    def export_docx(
        title: str,
        content: str,
        logo_path: str = None
    ) -> Tuple[BytesIO, str]:
        """Exporta documento para Word"""
        doc = Document()

        # Adicionar logo se fornecido
        if logo_path:
            try:
                doc.add_picture(logo_path, width=Inches(1.5))
            except Exception as e:
                print(f"Erro ao inserir logo: {e}")

        # Título
        title_para = doc.add_paragraph(title)
        title_para.style = 'Heading 1'
        for run in title_para.runs:
            run.font.size = Pt(24)
            run.font.color.rgb = RGBColor(31, 41, 55)

        # Conteúdo
        paragraphs = content.split('\n')
        for para in paragraphs:
            if para.strip():
                p = doc.add_paragraph(para)
                p.paragraph_format.line_spacing = 1.5
            else:
                doc.add_paragraph()

        # Rodapé
        footer_text = f"\n\nGerado em {datetime.now().strftime('%d/%m/%Y às %H:%M:%S')}"
        doc.add_paragraph(footer_text)

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        return buffer, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    @staticmethod
    def export_xlsx(
        title: str,
        content: str,
        data_filled: dict = None
    ) -> Tuple[BytesIO, str]:
        """Exporta documento para Excel"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Documento"

        # Header
        header_fill = PatternFill(start_color="1f2937", end_color="1f2937", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)

        ws['A1'] = title
        ws['A1'].font = Font(bold=True, size=14, color="1f2937")
        ws.row_dimensions[1].height = 25

        # Conteúdo
        row = 3
        paragraphs = content.split('\n')
        for para in paragraphs:
            if para.strip():
                ws[f'A{row}'] = para
                ws[f'A{row}'].alignment = Alignment(wrap_text=True, vertical='top')
                ws.row_dimensions[row].height = 30
                row += 1

        # Dados preenchidos se houver
        if data_filled:
            row += 2
            ws[f'A{row}'] = "Dados Preenchidos:"
            ws[f'A{row}'].font = Font(bold=True)
            row += 1

            for key, value in data_filled.items():
                ws[f'A{row}'] = f"{key}:"
                ws[f'B{row}'] = str(value)
                row += 1

        # Largura das colunas
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 30

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    @staticmethod
    def export_csv(
        title: str,
        data_filled: dict = None
    ) -> Tuple[BytesIO, str]:
        """Exporta dados para CSV"""
        buffer = BytesIO()

        # Usar StringIO para escrever em memória
        from io import StringIO
        string_buffer = StringIO()

        writer = csv.writer(string_buffer)
        writer.writerow(["Campo", "Valor"])
        writer.writerow([title, ""])
        writer.writerow([])

        if data_filled:
            for key, value in data_filled.items():
                writer.writerow([key, str(value)])

        # Converter para bytes
        buffer.write(string_buffer.getvalue().encode('utf-8-sig'))
        buffer.seek(0)

        return buffer, "text/csv"

    @staticmethod
    def export_pptx(
        title: str,
        content: str,
        logo_path: str = None
    ) -> Tuple[BytesIO, str]:
        """Exporta documento para PowerPoint"""
        prs = Presentation()
        prs.slide_width = PPTInches(10)
        prs.slide_height = PPTInches(7.5)

        # Slide 1: Título
        slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

        # Adicionar logo se fornecido
        if logo_path:
            try:
                left = PPTInches(0.5)
                top = PPTInches(0.5)
                height = PPTInches(0.8)
                slide1.shapes.add_picture(logo_path, left, top, height=height)
            except Exception as e:
                print(f"Erro ao inserir logo: {e}")

        # Título
        title_box = slide1.shapes.add_textbox(PPTInches(0.5), PPTInches(1.5), PPTInches(9), PPTInches(2))
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.word_wrap = True

        for paragraph in title_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = PPTPt(44)
                run.font.bold = True
                run.font.color.rgb = (31, 41, 55)

        # Slide 2: Conteúdo
        slide2 = prs.slides.add_slide(prs.slide_layouts[6])

        content_box = slide2.shapes.add_textbox(PPTInches(0.5), PPTInches(0.5), PPTInches(9), PPTInches(6.5))
        content_frame = content_box.text_frame
        content_frame.word_wrap = True

        paragraphs = content.split('\n')
        for i, para in enumerate(paragraphs):
            if i == 0:
                content_frame.text = para
            else:
                p = content_frame.add_paragraph()
                p.text = para

            if i < len(paragraphs) - 1:
                p = content_frame.add_paragraph()

        buffer = BytesIO()
        prs.save(buffer)
        buffer.seek(0)

        return buffer, "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    @staticmethod
    def export(
        format: str,
        title: str,
        content: str,
        logo_path: str = None,
        data_filled: dict = None
    ) -> Tuple[BytesIO, str]:
        """Função principal de exportação"""

        format_lower = format.lower()

        if format_lower == "pdf":
            return ExportService.export_pdf(title, content, logo_path)
        elif format_lower == "docx":
            return ExportService.export_docx(title, content, logo_path)
        elif format_lower == "xlsx":
            return ExportService.export_xlsx(title, content, data_filled)
        elif format_lower == "csv":
            return ExportService.export_csv(title, data_filled)
        elif format_lower == "pptx":
            return ExportService.export_pptx(title, content, logo_path)
        else:
            raise ValueError(f"Formato não suportado: {format}")
